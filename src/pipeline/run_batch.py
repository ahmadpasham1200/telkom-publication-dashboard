"""
Pipeline BATCH: scrape Google Scholar untuk BANYAK dosen sekaligus.

File ini tidak melakukan scraping maupun loading apa pun. Ia hanya
membaca daftar dosen dari Excel, memanggil run_pipeline.main() satu
per satu, dan mencatat hasil tiap dosen ke file progress CSV supaya
batch bisa dihentikan lalu dilanjutkan tanpa mengulang yang sudah
beres.

Alur
----
    1. Baca daftar dosen dari Excel (--file). Kolom `scholar_id`
       wajib ada; baris dengan scholar_id kosong tidak bisa
       diproses.
    2. Baris tanpa scholar_id dicatat DILEWATI di file progress
       dan TIDAK PERNAH dikirim ke scraper.
    3. Terapkan --mulai (potong baris awal) dan --lanjut (lewati
       yang statusnya sudah terminal di progress).
    4. Panggil run_pipeline.main(["--id", <uid>, ...]) per dosen.
       Setiap panggilan = 1 browser Playwright + 1 koneksi DB, dan
       tiap dosen ATOMIC: loader commit raw dulu lalu core, jadi
       kegagalan satu dosen tidak mempengaruhi dosen lain.
    5. Hasil tiap dosen langsung di-append + flush ke progress CSV.
    6. Ringkasan: total, dilewati, sukses, gagal, plus daftar
       dosen yang GAGAL saja.

Kenapa file progress di-flush setiap baris
------------------------------------------
Batch ini bisa jalan puluhan menit dan sering dihentikan pakai
Ctrl+C. Kalau baris progress ditahan di buffer, semua yang sudah
jalan hilang begitu proses mati - persis yang membuat --lanjut
tidak berguna. Karena itu setiap record ditulis lalu f.flush().

Exit code (main() mengembalikan int)
-------------------------------------
    0   = semua dosen yang diproses berhasil (atau tidak ada yang
          perlu diproses).
    1   = ada minimal satu dosen gagal (GAGAL_LOAD / GAGAL_SCRAPE /
          GAGAL).
    2   = gagal membaca daftar (file tidak ada / kolom scholar_id
          tidak ada).
    130 = KeyboardInterrupt; ringkasan tetap dicetak.

Exit code run_pipeline yang dibaca: 0 -> OK, 1 -> GAGAL_LOAD,
2 -> GAGAL_SCRAPE. Exit code 130 berarti interupsi: batch berhenti
rapi TANPA menulis record untuk dosen itu. Exception tak terduga
yang bocor dari run_pipeline.main -> GAGAL. SystemExit (dilempar
argparse lewat sys.exit, dan BUKAN Exception) juga ditangkap per
dosen di loop - lihat blok `except SystemExit` - supaya satu dosen
yang gagal parse tidak mematikan seluruh batch. Status DILEWATI
dipakai untuk baris tanpa profil.

Cara menjalankan
----------------
    python src/pipeline/run_batch.py --file "data dosen/Data dosen.xlsx"
    python src/pipeline/run_batch.py --file ... --lanjut --jeda 10
    python src/pipeline/run_batch.py --file ... --mulai 50 --ulang

Catatan: `--batch` di CLI ini BUKAN nama fitur batch runner; itu
batas klik "Tampilkan lainnya" yang diteruskan ke run_pipeline
(lihat run_pipeline.py). Nama "batch" untuk batch runner ada di
nama file ini, bukan di nama argumen.

Butuh paket: openpyxl (baca Excel), plus seluruh dependensi
run_pipeline (playwright, beautifulsoup4, sqlalchemy, psycopg,
python-dotenv).
"""

import argparse
import csv
import sys
import time
from datetime import datetime
from pathlib import Path

# --- Bootstrap import -------------------------------------------------
# Supaya file ini jalan baik sebagai skrip
# (`python src/pipeline/run_batch.py`) maupun sebagai modul
# (`python -m src.pipeline.run_batch`), keduanya dari root proyek.
_ROOT = Path(__file__).resolve().parents[2]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

# Windows default (cp1252) bisa crash saat print nama dosen
# beraksen atau pesan error dari psycopg.
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")  # type: ignore[attr-defined]
    except (AttributeError, ValueError):
        pass

from openpyxl import load_workbook  # noqa: E402

from src.pipeline import run_pipeline  # noqa: E402

# ============================================================
# KONSTANTA PROGRESS
# ============================================================

# Header file progress. Jangan diubah seenaknya: file yang sudah
# ada di lapangan harus tetap terbaca oleh baca_progress().
KOLOM_PROGRESS = ["scholar_id", "nama", "status", "exit_code", "waktu"]

STATUS_OK = "OK"
STATUS_GAGAL_LOAD = "GAGAL_LOAD"
STATUS_GAGAL_SCRAPE = "GAGAL_SCRAPE"
STATUS_DILEWATI = "DILEWATI"
STATUS_GAGAL = "GAGAL"

# Status yang berarti baris ini sudah selesai dan --lanjut boleh
# melewatkannya. GAGAL sengaja TIDAK ada di sini: kegagalan tak
# terduga harus dicoba ulang oleh --lanjut.
STATUS_TERMINAL = frozenset({
    STATUS_OK, STATUS_GAGAL_LOAD, STATUS_GAGAL_SCRAPE, STATUS_DILEWATI,
})

DEFAULT_PROGRESS = "data dosen/progress_batch.csv"

# Pemetaan exit code run_pipeline.main() -> status progress.
# 130 sengaja tidak ada di sini: artinya interupsi, dan interupsi
# menghentikan seluruh batch, bukan mencatat satu dosen gagal.
_PETA_STATUS = {
    0: STATUS_OK,
    1: STATUS_GAGAL_LOAD,
    2: STATUS_GAGAL_SCRAPE,
}


class KolomScholarIdHilang(SystemExit, argparse.ArgumentTypeError):
    """Kolom `scholar_id` tidak ditemukan di file Excel.

    Turunan ganda SystemExit + argparse.ArgumentTypeError supaya
    penanganannya sah di dua konteks: dari dalam main() ia bisa
    ditangkap sebagai error keluaran (atau lolos ke interpreter
    sebagai exit code 1 dengan pesan ke stderr), dan dari konteks
    argparse ia terbaca seperti error argumen biasa. Pesannya selalu
    memuat teks `scholar_id` plus kolom yang harus ditambahkan.
    """


# ============================================================
# BACA DAFTAR DARI EXCEL
# ============================================================

def _teks_nilai(nilai):
    """Normalisasi sel Excel -> string; kosong (None/putih) -> ""."""
    if nilai is None:
        return ""
    if isinstance(nilai, float) and nilai.is_integer():
        # 123.0 jangan berubah jadi "123.0".
        nilai = int(nilai)
    return str(nilai).strip()


def baca_daftar(file_path: str) -> list:
    """Baca daftar dosen dari Excel -> list dict per baris DATA.

    Tiap dict berisi `no`, `nama`, `scholar_id` (string hasil strip;
    sel kosong/None/putih -> ""). Kolom `scholar_id` dideteksi
    case-insensitive terhadap baris pertama (header). Baris yang
    seluruhnya kosong dilewati.

    Melempar KolomScholarIdHilang kalau kolom `scholar_id` tidak
    ada, dan OSError kalau file-nya tidak bisa dibuka.
    """
    try:
        wb = load_workbook(filename=str(file_path), read_only=True,
                           data_only=True)
    except OSError:
        raise
    except Exception as exc:
        # File bukan Excel valid (atau corrupt): laporkan sebagai
        # kegagalan membaca daftar, bukan traceback panjang.
        raise OSError(
            "File {} tidak bisa dibaca sebagai Excel: {}".format(
                file_path, exc)) from exc

    try:
        if "Sheet1" in wb.sheetnames:
            ws = wb["Sheet1"]
        else:
            ws = wb[wb.sheetnames[0]]

        baris_iter = ws.iter_rows(values_only=True)
        header = next(baris_iter, None)
        if header is None:
            raise KolomScholarIdHilang(
                "File {} tidak punya baris header, jadi kolom `scholar_id` "
                "tidak ditemukan. Tambahkan kolom `scholar_id` (ID Google "
                "Scholar dari bagian ?user= URL profil) pada baris header "
                "file Excel itu.".format(file_path))

        judul = [("" if c is None else str(c).strip().lower())
                 for c in header]

        if "scholar_id" not in judul:
            raise KolomScholarIdHilang(
                "Kolom `scholar_id` tidak ditemukan di {}. Kolom yang ada: "
                "{}. Tambahkan kolom `scholar_id` (diisi ID Google Scholar "
                "dari bagian ?user= URL profil) pada file Excel itu."
                .format(file_path, ", ".join(
                    [str(c).strip() for c in header if c is not None]) or
                    "(tidak ada)"))

        # `no` dan `nama` dicari case-insensitive juga; kalau tidak
        # ada, jatuh ke posisi kolom pertama dan kedua.
        idx_scholar = judul.index("scholar_id")
        idx_no = judul.index("no") if "no" in judul else 0
        idx_nama = judul.index("nama") if "nama" in judul else 1

        def _sel(baris, idx):
            if idx >= len(baris):
                return ""
            return baris[idx]

        hasil = []
        for baris in baris_iter:
            # Baris kosong total (sering muncul di Excel) dilewati,
            # bukan dihitung sebagai dosen tanpa profil.
            if all(c is None or str(c).strip() == "" for c in baris):
                continue
            hasil.append({
                "no": _sel(baris, idx_no),
                "nama": _teks_nilai(_sel(baris, idx_nama)),
                "scholar_id": _teks_nilai(_sel(baris, idx_scholar)),
            })
        return hasil
    finally:
        wb.close()


# ============================================================
# FILE PROGRESS (CSV)
# ============================================================

def _pecah_baris(baris_mentah, delimiter):
    """Pecah satu baris CSV dengan csv.reader -> list field.

    csv.reader dipakai (bukan str.split) karena ia menghormati tanda
    kutip, jadi baris yang kebetulan mengandung delimiter di dalam
    kutip tidak ikut terbelah.
    """
    return next(csv.reader([baris_mentah], delimiter=delimiter))


def baca_progress(path: str) -> dict:
    """Baca CSV progress -> map `scholar_id` -> baris (dict).

    File yang belum ada atau masih kosong -> {}. Baris kembar
    (mis. banyak dosen DILEWATI berbagi scholar_id "") dipakai
    yang terakhir, karena kuncinya memang scholar_id.

    Delimiter dipilih PER BARIS, bukan per file. Kenapa: file
    progress pernah dibuka/di-save ulang oleh Excel dengan locale
    Indonesia (pemisah `;`), dan hasilnya file MIXED - header dan
    sebagian baris memakai `;`, baris lain (yang ditulis ulang oleh
    batch ini dengan tulis_progress) memakai `,`. Deteksi per-file
    (lihat header saja) GAGAL pada kasus ini: header `;` dipilih,
    lalu baris koma yang tidak mengandung `;` sama sekali terbaca
    sebagai SATU field utuh, kuncinya jadi kalimat lengkap, dan
    scholar_id asli tidak pernah muncul - --lanjut diam-diam
    mengulang dosen yang sudah jalan. Jangan hapus pemilihan per
    baris ini.

    Aturan per baris data:
      1. Nama field ditentukan header (delimiter header dipilih
         yang menghasilkan field terbanyak; seri -> koma, format
         penulis tulis_progress).
      2. Tiap baris dipecah dengan `;` dan dengan `,`; dipakai yang
         jumlah fieldnya SAMA dengan jumlah nama field. Kalau
         keduanya sama-sama cocok -> koma (format penulis). Baris
         koma yang dipecah `;` menghasilkan 1 field, jadi otomatis
         kalah hitungan - tidak perlu heuristik lain.
      3. Kalau tidak ada yang cocok -> baris rusak, DIABAIKAN
         (tidak dilempar, tidak masuk hasil).
      4. Baris yang seluruh field kosong (mis. `;;;;`) diabaikan.
    """
    p = Path(path)
    if not p.exists() or p.stat().st_size == 0:
        return {}
    hasil = {}
    with p.open("r", encoding="utf-8", newline="") as f:
        baris_header = f.readline()
        if not baris_header.strip():
            return {}

        header_koma = _pecah_baris(baris_header, ",")
        header_titik = _pecah_baris(baris_header, ";")
        if len(header_titik) > len(header_koma):
            nama_field = header_titik
        else:
            nama_field = header_koma
        jumlah_field = len(nama_field)

        for baris_mentah in f:
            pecah_koma = _pecah_baris(baris_mentah, ",")
            pecah_titik = _pecah_baris(baris_mentah, ";")

            if (len(pecah_titik) == jumlah_field
                    and len(pecah_koma) != jumlah_field):
                field = pecah_titik
            else:
                # Koma menang seri (format penulis) dan menang kalau
                # hanya koma yang cocok.
                field = pecah_koma

            if len(field) != jumlah_field:
                # Baris rusak: jumlah kolom tidak cocok header.
                continue
            if not any((v or "").strip() for v in field):
                # Baris kosong total (mis. ";;;;") diabaikan.
                continue

            baris = dict(zip(nama_field, field))
            kunci = baris.get("scholar_id")
            hasil["" if kunci is None else kunci] = baris
    return hasil


def tulis_progress(path: str, record: dict) -> None:
    """Append 1 baris ke CSV progress lalu FLUSH.

    Flush-nya wajib: batch bisa mati lewat Ctrl+C kapan saja, dan
    baris yang tertahan di buffer akan hilang - membuat --lanjut
    kehilangan pekerjaan yang sudah jalan. Header ditulis hanya
    kalau file belum ada/ masih kosong. `waktu` diisi otomatis
    ISO-8601 bila record tidak membawanya.
    """
    p = Path(path)
    if p.parent != Path(""):
        p.parent.mkdir(parents=True, exist_ok=True)

    perlu_header = (not p.exists()) or p.stat().st_size == 0
    with p.open("a", encoding="utf-8", newline="") as f:
        penulis = csv.DictWriter(f, fieldnames=KOLOM_PROGRESS,
                                 lineterminator="\n",
                                 extrasaction="ignore")
        if perlu_header:
            penulis.writeheader()

        baris = {}
        for kolom in KOLOM_PROGRESS:
            nilai = record.get(kolom, "")
            if nilai is None:
                nilai = ""
            baris[kolom] = nilai
        if not str(baris["waktu"]).strip():
            baris["waktu"] = datetime.now().isoformat(timespec="seconds")
        penulis.writerow(baris)
        f.flush()


# ============================================================
# PILIHAN OPSIONAL: BACA DB UNTUK MEMBEDAKAN LOADED vs SKIPPED
# ============================================================

def _statistik_dari_db(scholar_ids, url=None) -> dict:
    """Baca raw.scholar_scrape_result untuk run terakhir tiap scholar.

    Kembalikan map scholar_id -> dict {status_run, ada_baris}.
    `ada_baris` membedakan LOADED (envelope punya baris, core
    kemungkinan terisi) vs SKIPPED (envelope tanpa baris).

    Sifatnya penambah nilai: kalau database tidak bisa diakses,
    pemanggil menangkap exceptionnya dan melewatinya dengan
    peringatan. Fungsi ini tidak pernah dipanggil dari test
    (di-patch agar test tidak butuh PostgreSQL).
    """
    if not scholar_ids:
        return {}
    import sqlalchemy as sa
    from src.loading.postgres import connect

    sql = sa.text(
        "SELECT DISTINCT ON (scholar_id) "
        "       scholar_id, status, "
        "       jsonb_array_length("
        "           coalesce(raw_envelope->'data'->'rows', '[]'::jsonb)"
        "       ) AS n_rows "
        "FROM raw.scholar_scrape_result "
        "WHERE scholar_id IN :ids "
        "ORDER BY scholar_id, scraped_at DESC"
    ).bindparams(sa.bindparam("ids", expanding=True))

    conn = connect(url=url)
    try:
        hasil = {}
        for sid, status_run, n_rows in conn.execute(
                sql, {"ids": list(scholar_ids)}):
            hasil[sid] = {
                "status_run": status_run,
                "ada_baris": bool(n_rows),
            }
        return hasil
    finally:
        conn.close()


# ============================================================
# CLI
# ============================================================

def parse_args(argv=None):
    p = argparse.ArgumentParser(
        description="Scrape daftar dosen dari Excel, satu per satu, "
                    "lalu catat hasilnya ke file progress CSV.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "Exit code:\n"
            "  0   semua dosen yang diproses berhasil\n"
            "  1   ada dosen gagal (GAGAL_LOAD/GAGAL_SCRAPE/GAGAL)\n"
            "  2   gagal membaca daftar (file / kolom scholar_id)\n"
            "  130 dihentikan pengguna (Ctrl+C)\n"
            "\n"
            "Contoh:\n"
            "  python src/pipeline/run_batch.py --file \"data dosen/Data dosen.xlsx\"\n"
            "  python src/pipeline/run_batch.py --file ... --lanjut --jeda 10\n"
            "  python src/pipeline/run_batch.py --file ... --mulai 50 --ulang\n"
        ),
    )
    p.add_argument("--file", required=True,
                   help="Path file Excel daftar dosen (kolom `scholar_id` "
                        "wajib ada; baris kosong = DILEWATI).")
    p.add_argument("--jeda", type=float, default=8.0,
                   help="Jeda antar dosen dalam detik (default 8.0)")
    p.add_argument("--lanjut", action="store_true",
                   help="Lewati dosen yang statusnya di file progress sudah "
                        "terminal (OK, GAGAL_LOAD, GAGAL_SCRAPE, DILEWATI)")
    p.add_argument("--ulang", action="store_true",
                   help="Paksa ulang semua dosen, mengalahkan --lanjut")
    p.add_argument("--mulai", type=int, default=1,
                   help="Mulai dari baris data ke-N, 1-based dihitung dari "
                        "baris DATA pertama, bukan header (default 1)")
    p.add_argument("--progress", default=DEFAULT_PROGRESS,
                   help="Path file progress CSV (default: %(default)s)")
    # --- diteruskan ke run_pipeline.main() ---
    p.add_argument("--headed", action="store_true",
                   help="[diteruskan] Tampilkan browser")
    p.add_argument("--batch", type=int, default=None,
                   help="[diteruskan] Batas klik 'Tampilkan lainnya' di "
                        "run_pipeline (bukan jumlah dosen)")
    p.add_argument("--diam", action="store_true",
                   help="[diteruskan] Matikan progress scraper (selalu "
                        "dinyalakan untuk batch, flag ini dipertahankan "
                        "agar antarmuka tetap sama)")
    p.add_argument("--tampil", type=int, default=None,
                   help="[diteruskan] Jumlah baris tabel yang dicetak "
                        "run_pipeline per dosen (default 0 = tidak mencetak)")
    p.add_argument("--url", default=None,
                   help="[diteruskan] URL koneksi PostgreSQL lengkap")
    p.add_argument("--echo-sql", action="store_true",
                   help="[diteruskan] Cetak SQL yang dijalankan")
    return p.parse_args(argv)


def _argv_run_pipeline(args, uid: str) -> list:
    """Susun argv untuk run_pipeline.main() untuk satu dosen.

    Selalu membawa `--id` dan `--diam`; `--tampil 0` dipasang bila
    user tidak menyetel --tampil, supaya batch tidak dibanjiri tabel
    puluhan dosen.

    SEMUA argumen nilai memakai bentuk `--nama=nilai`, bukan dua
    token terpisah. Ini perbaikan bug, bukan selera gaya: ada
    scholar_id yang berawalan minus (mis. `-zWuOpYAAAAJ`; pola
    `^[A-Za-z0-9_-]{7}AAAAJ$` memang mengizinkan `-`). Dengan dua
    token, `--id -zWuOpYAAAAJ` dibaca argparse sebagai opsi, bukan
    nilai -> `error: argument --id: expected one argument` ->
    sys.exit(2) -> batch mati tepat di dosen itu. Bentuk
    `--id=-zWuOpYAAAAJ` aman karena tokennya diawali `--id=`,
    bukan `-`. Konsisten diterapkan ke `--id`, `--batch`,
    `--tampil`, dan `--url`; `--headed`/`--diam`/`--echo-sql`
    memang flag tanpa nilai.
    """
    argv = ["--id={}".format(uid), "--diam"]
    if args.headed:
        argv.append("--headed")
    if args.batch is not None:
        argv.append("--batch={}".format(args.batch))
    argv.append("--tampil={}".format(
        0 if args.tampil is None else args.tampil))
    if args.url:
        argv.append("--url={}".format(args.url))
    if args.echo_sql:
        argv.append("--echo-sql")
    return argv


def _garis(judul):
    print("\n" + "=" * 60)
    print(judul)
    print("=" * 60)


def _catat_keluaran_terakhir(nama, uid, hasil, exit_code):
    print("  -> {} ({}): {} exit={}".format(
        nama, uid, hasil, exit_code if exit_code != "" else "-"))


def main(argv=None):
    args = parse_args(argv)
    if args.mulai < 1:
        args.mulai = 1
    if args.jeda < 0:
        args.jeda = 0.0

    _garis("BATCH SCRAPE -> POSTGRESQL")
    print("  File Excel : {}".format(args.file))
    print("  Progress   : {}".format(args.progress))
    print("  Jeda       : {} detik".format(args.jeda))
    print("  Mulai      : baris data ke-{}".format(args.mulai))
    print("  Mode       : {}".format(
        "lanjut (lewati yang sudah terminal)" if args.lanjut and not args.ulang
        else "ulang (paksa semua)" if args.ulang
        else "dari awal"))

    # ---------- 1. BACA DAFTAR ----------
    try:
        daftar = baca_daftar(args.file)
    except KolomScholarIdHilang as exc:
        # Flush stdout dulu: kalau output di-pipe, stderr tanpa flush
        # bisa muncul sebelum banner di atas.
        sys.stdout.flush()
        print("GAGAL membaca daftar: {}".format(exc), file=sys.stderr)
        return 2
    except OSError as exc:
        sys.stdout.flush()
        print("GAGAL membaca daftar: {}".format(exc), file=sys.stderr)
        return 2

    progress = baca_progress(args.progress)
    total = len(daftar)
    tanpa_profil = [r for r in daftar if not r["scholar_id"]]
    print("  Baris data : {}".format(total))
    print("  Tanpa profil Scholar: {}".format(len(tanpa_profil)))

    # ---------- 2. DILEWATI (tanpa profil) ----------
    # Sengaja dilakukan SEBELUM filter --mulai: semua dosen tanpa
    # profil tetap terdokumentasi di progress, apa pun --mulai.
    for r in tanpa_profil:
        sudah = any(
            rec.get("status") == STATUS_DILEWATI and
            rec.get("nama") == r["nama"]
            for rec in progress.values()
        )
        if sudah:
            continue
        record = {
            "scholar_id": "",
            "nama": r["nama"],
            "status": STATUS_DILEWATI,
            "exit_code": "",
            "waktu": "",
        }
        tulis_progress(args.progress, record)
        progress[""] = record
        print("  DILEWATI  : {} (tidak punya profil Scholar)".format(
            r["nama"]))

    # ---------- 3 & 4. FILTER --mulai / --lanjut ----------
    kandidat = []
    dilewati_lanjut = 0
    for nomor, r in enumerate(daftar, start=1):  # nomor baris DATA, 1-based
        if not r["scholar_id"]:
            continue  # sudah dicatat DILEWATI di atas
        if nomor < args.mulai:
            continue
        if args.lanjut and not args.ulang:
            rec = progress.get(r["scholar_id"])
            if rec and rec.get("status") in STATUS_TERMINAL:
                dilewati_lanjut += 1
                print("  LEWAT     : {} ({}) sudah {}".format(
                    r["nama"], r["scholar_id"], rec.get("status")))
                continue
        kandidat.append(r)

    print("  Akan diproses: {}".format(len(kandidat)))
    if dilewati_lanjut:
        print("  Dilewati (--lanjut): {}".format(dilewati_lanjut))

    # ---------- 5. LOOP PER DOSEN ----------
    ringkasan = []  # (record, nama, scholar_id)
    dipanggil = []  # scholar_id yang benar-benar masuk run_pipeline
    terputus = False
    n = len(kandidat)

    try:
        for idx, r in enumerate(kandidat):
            uid = r["scholar_id"]
            argv_anak = _argv_run_pipeline(args, uid)
            print("\n[{}/{}] {} ({})".format(idx + 1, n, r["nama"], uid))
            dipanggil.append(uid)

            try:
                exit_code = run_pipeline.main(argv_anak)
            except KeyboardInterrupt:
                # Interupsi di dalam run_pipeline: berhenti rapi,
                # jangan catat apa pun (dosen ini bisa jadi setengah
                # jalan; biar --lanjut mengulangnya nanti).
                terputus = True
                break
            except SystemExit as exc:
                # SystemExit harus ditangkap DI TINGKAT INI, sejajar
                # Exception, karena ia BaseException - BUKAN Exception.
                # argparse memanggil sys.exit() (melempar SystemExit)
                # saat gagal parse, jadi kalau hanya `except
                # Exception` yang ada, SystemExit lolos dari loop,
                # lolos dari try luar, dan mematikan SELURUH batch di
                # tengah jalan (persis yang terjadi saat satu
                # scholar_id membuat run_pipeline error parse).
                # Kontraknya dipertahankan: satu dosen bermasalah
                # tidak boleh mematikan dosen berikutnya.
                if exc.code == 130:
                    # Sama seperti KeyboardInterrupt: ini permintaan
                    # berhenti, bukan kegagalan satu dosen.
                    terputus = True
                    break
                if isinstance(exc.code, int):
                    # Kode keluaran angka -> perlakukan seperti
                    # run_pipeline.main() yang mengembalikan kode itu,
                    # ikut jalur _PETA_STATUS seperti biasa.
                    exit_code = exc.code
                    status = _PETA_STATUS.get(exit_code, STATUS_GAGAL)
                    exit_code_simpan = exit_code
                    print("  SystemExit(code={}) dari run_pipeline "
                          "ditangkap; batch tetap lanjut.".format(
                              exit_code))
                else:
                    # sys.exit() tanpa kode angka (None) atau kode
                    # aneh -> tidak ada exit code yang bisa dipetakan,
                    # perlakukan seperti kegagalan tak terduga.
                    status = STATUS_GAGAL
                    exit_code_simpan = ""
                    print("  SystemExit tanpa kode angka ({!r}) dari "
                          "run_pipeline; dicatat GAGAL.".format(
                              exc.code))
            except Exception as exc:
                # Exception tak terduga TIDAK boleh mematikan batch:
                # satu dosen bermasalah, dosen berikutnya tetap jalan.
                status = STATUS_GAGAL
                exit_code_simpan = ""
                print("  GAGAL tak terduga: {}: {}".format(
                    type(exc).__name__, exc))
            else:
                if exit_code == 130:
                    # run_pipeline menangkap KeyboardInterrupt sendiri.
                    terputus = True
                    break
                status = _PETA_STATUS.get(exit_code, STATUS_GAGAL)
                exit_code_simpan = exit_code

            record = {
                "scholar_id": uid,
                "nama": r["nama"],
                "status": status,
                "exit_code": exit_code_simpan,
                "waktu": "",
            }
            tulis_progress(args.progress, record)
            progress[uid] = record
            ringkasan.append((record, r["nama"], uid))
            _catat_keluaran_terakhir(r["nama"], uid, status,
                                     exit_code_simpan)

            # Jeda antar iterasi; baris terakhir tidak perlu.
            if idx < n - 1:
                time.sleep(args.jeda)
    except KeyboardInterrupt:
        terputus = True

    # ---------- RINGKASAN ----------
    _garis("RINGKASAN BATCH")
    sukses = [x for x in ringkasan if x[0]["status"] == STATUS_OK]
    gagal = [x for x in ringkasan
             if x[0]["status"] != STATUS_OK]

    print("  Total baris data      : {}".format(total))
    print("  Dilewati (tanpa profil): {}".format(len(tanpa_profil)))
    print("  Dilewati (--lanjut)    : {}".format(dilewati_lanjut))
    print("  Dipanggil scraper      : {}".format(len(dipanggil)))
    print("  Sukses (OK)            : {}".format(len(sukses)))
    print("  Gagal                  : {}".format(len(gagal)))

    # Penambah nilai: kalau DB bisa dibaca, tunjukkan mana yang
    # LOADED vs SKIPPED. Gagal membaca DB jangan pernah crash.
    if sukses:
        try:
            db = _statistik_dari_db(
                [sid for _, _, sid in sukses], url=args.url)
        except Exception as exc:
            print("  (peringatan) Tidak bisa membaca "
                  "raw.scholar_scrape_result: {}: {}".format(
                      type(exc).__name__, exc))
            print("           Ringkasan LOADED/SKIPPED dilewati.")
        else:
            if db:
                loaded = [sid for sid, v in db.items() if v["ada_baris"]]
                skipped = [sid for sid, v in db.items()
                           if not v["ada_baris"]]
                print()
                print("  Dari raw.scholar_scrape_result (run terakhir):")
                print("    punya baris (kemungkinan LOADED) : {}"
                      .format(len(loaded)))
                print("    tanpa baris (kemungkinan SKIPPED): {}"
                      .format(len(skipped)))
                for sid in skipped:
                    print("      - {} (SKIPPED, envelope tanpa baris)"
                          .format(sid))
            else:
                print("  (peringatan) Tidak ada record di "
                      "raw.scholar_scrape_result untuk dosen yang OK.")

    if gagal:
        print()
        print("  DOSEN YANG GAGAL:")
        for record, nama, uid in gagal:
            print("    - {} ({}): {} exit={}".format(
                nama, uid, record["status"],
                record["exit_code"] if record["exit_code"] != "" else "-"))

    if terputus:
        print()
        print("  Dihentikan pengguna. Ringkasan di atas tetap valid; "
              "jalankan lagi dengan --lanjut untuk melanjutkan.")
        return 130

    if gagal:
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())

